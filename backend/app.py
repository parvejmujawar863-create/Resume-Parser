
import re
from pathlib import Path

from flask import Flask, jsonify, request
from flask_cors import CORS
from pypdf import PdfReader
from docx import Document
from werkzeug.utils import secure_filename

app = Flask(__name__)
CORS(app)

app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

SKILLS = {
    "python", "java", "javascript", "html", "css",
    "react", "node.js", "flask", "django",
    "sql", "mysql", "mongodb", "git", "github",
    "c", "c++", "php", "excel", "communication",
    "problem solving", "data analysis", "machine learning",
    "rest api", "linux", "cybersecurity", "aws",
    "typescript", "bootstrap", "tailwind"
}


def extract_text(file_path, extension):
    """Extract text from a supported resume file."""
    if extension == ".pdf":
        reader = PdfReader(file_path)

        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")

        pages = reader.pages[:30]
        return "\n".join(page.extract_text() or "" for page in pages)

    if extension == ".docx":
        document = Document(file_path)
        paragraphs = [p.text for p in document.paragraphs]
        return "\n".join(paragraphs)

    if extension == ".txt":
        return file_path.read_bytes().decode("utf-8-sig")

    raise ValueError("Unsupported file type.")


def find_skills(text):
    """Find known skills using case-insensitive whole-term matching."""
    found = set()
    normalized_text = text.lower()

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"

        if re.search(pattern, normalized_text):
            found.add(skill)

    return sorted(found)


@app.get("/")
def home():
    return jsonify({
        "message": "ResumeIQ API is running.",
        "status": "ok"
    })


@app.get("/api/health")
def health():
    return jsonify({"status": "healthy"})


@app.post("/api/analyze")
def analyze_resume():
    if "resume" not in request.files:
        return jsonify({"error": "Please upload a resume file."}), 400

    resume = request.files["resume"]

    if not resume.filename:
        return jsonify({"error": "No file selected."}), 400

    safe_name = secure_filename(resume.filename)
    extension = Path(safe_name).suffix.lower()

    if not safe_name or extension not in ALLOWED_EXTENSIONS:
        return jsonify({
            "error": "Upload a PDF, DOCX, or TXT file."
        }), 400

    job_description = request.form.get("job_description", "").strip()

    if len(job_description) > 15000:
        return jsonify({
            "error": "Job description must be 15,000 characters or less."
        }), 400

    try:
        # Process in memory: the uploaded file is not saved to disk.
        from io import BytesIO

        file_data = resume.read()

        if not file_data:
            return jsonify({"error": "The uploaded file is empty."}), 400

        if extension == ".pdf":
            reader = PdfReader(BytesIO(file_data))

            if reader.is_encrypted:
                return jsonify({
                    "error": "Password-protected PDFs are not supported."
                }), 400

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages[:30]
            )

        elif extension == ".docx":
            document = Document(BytesIO(file_data))
            text = "\n".join(p.text for p in document.paragraphs)

        else:
            text = file_data.decode("utf-8-sig")

        text = text.strip()

        if not text:
            return jsonify({
                "error": "No readable text found. Scanned PDFs may need OCR."
            }), 422

        resume_skills = find_skills(text)
        job_skills = find_skills(job_description)

        matched = sorted(set(resume_skills) & set(job_skills))
        missing = sorted(set(job_skills) - set(resume_skills))

        match_score = (
            round(len(matched) / len(job_skills) * 100)
            if job_skills else None
        )

        return jsonify({
            "filename": safe_name,
            "character_count": len(text),
            "skills": resume_skills,
            "skills_count": len(resume_skills),
            "job_skills": job_skills,
            "matched_skills": matched,
            "missing_skills": missing,
            "match_score": match_score,
            "message": (
                "Analysis complete. The match score is a simple "
                "keyword overlap, not an assessment of hiring potential."
            )
        })

    except (ValueError, UnicodeDecodeError) as exc:
        return jsonify({"error": str(exc)}), 400

    except Exception:
        app.logger.exception("Resume analysis failed")
        return jsonify({
            "error": "Could not process this file. Please try another supported resume."
        }), 500


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "File exceeds the 5 MB upload limit."}), 413


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
